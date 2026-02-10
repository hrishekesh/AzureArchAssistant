from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass
class RequirementDocument:
    title: str
    content: str
    assumptions: list[str]


def _normalize_list(value: Any) -> list[str]:
    if not value:
        return []
    if isinstance(value, str):
        return [value]
    return [str(item).strip() for item in value if str(item).strip()]


def _is_unclear(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        normalized = value.strip().lower()
        return normalized == "" or normalized in {"tbd", "n/a", "unknown", "not provided"}
    if isinstance(value, list):
        return len(_normalize_list(value)) == 0
    return False


def _build_assumptions(payload: dict[str, Any]) -> list[str]:
    assumptions = _normalize_list(payload.get("assumptions"))

    if _is_unclear(payload.get("functional_requirements")):
        assumptions.append(
            "Functional requirements are incomplete. The solution should prioritize extensibility to accommodate additional requirements."
        )

    if _is_unclear(payload.get("non_functional_requirements")):
        assumptions.append(
            "Non-functional requirements are incomplete. Baseline quality attributes assume security, reliability, and observability best practices."
        )

    if _is_unclear(payload.get("architecture_overview")):
        assumptions.append(
            "Architecture overview details were not fully provided. A layered architecture with clear API, domain, and persistence boundaries is assumed."
        )

    if _is_unclear(payload.get("high_level_solution_summary")):
        assumptions.append(
            "High-level solution summary was not fully provided. The recommendation assumes cloud-native deployment and incremental delivery."
        )

    if _is_unclear(payload.get("software_components")):
        assumptions.append(
            "Software component definitions are incomplete. Capability mapping may include placeholder components pending clarification."
        )

    # Preserve order while removing duplicates
    deduped: list[str] = []
    for assumption in assumptions:
        if assumption not in deduped:
            deduped.append(assumption)
    return deduped


def _build_capability_matrix(payload: dict[str, Any]) -> list[tuple[str, list[str]]]:
    raw_matrix = payload.get("capabilities")
    matrix: list[tuple[str, list[str]]] = []

    if isinstance(raw_matrix, list) and raw_matrix:
        for item in raw_matrix:
            if not isinstance(item, dict):
                continue
            capability = str(item.get("capability", "")).strip()
            components = _normalize_list(item.get("components"))
            if capability:
                matrix.append((capability, components or ["TBD Component"]))

    if matrix:
        return matrix

    components = _normalize_list(payload.get("software_components"))
    fallback_components = components or ["TBD Component"]
    for req in _normalize_list(payload.get("functional_requirements")):
        matrix.append((req, fallback_components))

    if not matrix:
        matrix.append(("Unspecified Capability", fallback_components))

    return matrix


def assemble_document(payload: dict[str, Any]) -> RequirementDocument:
    title = str(payload.get("project_name", "Architecture Requirements Document")).strip() or "Architecture Requirements Document"

    functional_requirements = _normalize_list(payload.get("functional_requirements"))
    non_functional_requirements = _normalize_list(payload.get("non_functional_requirements"))
    architecture_overview = str(payload.get("architecture_overview", "")).strip() or "Not provided."
    high_level_solution_summary = str(payload.get("high_level_solution_summary", "")).strip() or "Not provided."

    assumptions = _build_assumptions(payload)
    capability_matrix = _build_capability_matrix(payload)

    lines: list[str] = [f"# {title}", ""]

    lines.extend(["## Functional Requirements", ""])
    if functional_requirements:
        lines.extend([f"- {item}" for item in functional_requirements])
    else:
        lines.append("- Not provided.")

    lines.extend(["", "## Non-Functional Requirements", ""])
    if non_functional_requirements:
        lines.extend([f"- {item}" for item in non_functional_requirements])
    else:
        lines.append("- Not provided.")

    lines.extend(["", "## Architecture Overview", "", architecture_overview])
    lines.extend(["", "## High-Level Solution Summary", "", high_level_solution_summary])

    lines.extend(["", "## Capability Matrix", "", "| Capability | Software Components |", "|---|---|"])
    for capability, components in capability_matrix:
        lines.append(f"| {capability} | {', '.join(components)} |")

    lines.extend(["", "## Assumptions", ""])
    if assumptions:
        lines.extend([f"- {item}" for item in assumptions])
    else:
        lines.append("- No assumptions required.")

    return RequirementDocument(title=title, content="\n".join(lines), assumptions=assumptions)


def export_document(document: RequirementDocument, output_path: str | Path) -> Path:
    destination = Path(output_path)
    suffix = destination.suffix.lower()

    if suffix == ".docx":
        return _export_docx(document, destination)

    if suffix == ".md":
        destination.write_text(document.content, encoding="utf-8")
        return destination

    raise ValueError("Unsupported output format. Use .docx or .md")


def _export_docx(document: RequirementDocument, destination: Path) -> Path:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX export requires python-docx. Install it with: pip install python-docx"
        ) from exc

    doc = Document()
    for line in document.content.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return destination
